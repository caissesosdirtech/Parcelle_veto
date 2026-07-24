from rest_framework import viewsets
from django.shortcuts import render, redirect, get_object_or_404

from .models import Vente, LigneVente
from .serializers import VenteSerializer

from django.http import HttpResponse, JsonResponse
from reportlab.pdfgen import canvas

from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.utils.timezone import now

from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST

from pharmacie.models import Medicament
from clients.models import Client
from consultations.models import Consultation

from django.db import transaction
import json

# ===================== EXPORTS (Excel / PDF) =====================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


# ===================== DRF =====================

class VenteViewSet(viewsets.ModelViewSet):
    queryset = Vente.objects.all().order_by('-id')
    serializer_class = VenteSerializer


# ===================== VUES WEB (HTML) =====================

def ventes_list(request):
    today = now()
    ventes = Vente.objects.all().order_by('-date')

    search = request.GET.get('search', '').strip()
    type_filtre = request.GET.get('type', '').strip()
    date_filtre = request.GET.get('date', '').strip()

    if search:
        ventes = ventes.filter(
            Q(client__nom__icontains=search) |
            Q(ordonnance__consultation__client__nom__icontains=search) |
            Q(ordonnance__consultation__animal__nom__icontains=search) |
            Q(id__icontains=search)
        ).distinct()

    if type_filtre == 'ordonnance':
        ventes = ventes.filter(ordonnance__isnull=False)
    elif type_filtre == 'direct':
        ventes = ventes.filter(ordonnance__isnull=True)

    if date_filtre:
        ventes = ventes.filter(date__date=date_filtre)

    total_ventes = ventes.aggregate(Sum('total'))['total__sum'] or 0

    ventes_aujourdhui = Vente.objects.filter(
        date__year=today.year,
        date__month=today.month,
        date__day=today.day
    ).aggregate(Sum('total'))['total__sum'] or 0

    nb_ventes = ventes.count()
    ticket_moyen = round(total_ventes / nb_ventes, 0) if nb_ventes > 0 else 0

    consultations_jour = Consultation.objects.filter(
        date__year=today.year,
        date__month=today.month,
        date__day=today.day
    ).count()

    return render(request, "ventes/liste_ventes.html", {
        "ventes": ventes,
        "total_ventes": total_ventes,
        "ventes_aujourdhui": ventes_aujourdhui,
        "nb_ventes": nb_ventes,
        "ticket_moyen": ticket_moyen,
        "consultations_jour": consultations_jour,
    })


def vente_pdf(request, id):
    vente = Vente.objects.get(id=id)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="vente.pdf"'

    p = canvas.Canvas(response)

    p.setFont("Helvetica-Bold", 16)
    p.drawString(120, 800, "CABINET VÉTÉRINAIRE PARCELLE VETO")

    p.setFont("Helvetica", 11)
    p.drawString(50, 770, f"Client : {vente.ordonnance.consultation.client.nom}")
    p.drawString(50, 750, f"Animal : {vente.ordonnance.consultation.animal.nom}")
    p.drawString(50, 730, f"Date : {vente.date}")

    y = 690
    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "Médicaments :")

    y -= 20

    p.setFont("Helvetica", 10)

    for ligne in vente.lignes.all():
        p.drawString(
            60,
            y,
            f"{ligne.medicament.nom} | Qté:{ligne.quantite} | PU:{ligne.prix_unitaire} | Total:{ligne.montant_total}"
        )
        y -= 20

        if y < 100:
            p.showPage()
            y = 800

    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y - 30, f"TOTAL : {vente.total} FCFA")

    p.setFont("Helvetica", 9)
    p.drawString(150, 40, "Merci pour votre confiance - Parcelle Veto")

    p.save()
    return response


def vente_detail(request, vente_id):
    vente = get_object_or_404(
        Vente.objects.prefetch_related("lignes__medicament__catalogue"),
        id=vente_id
    )
    return render(
        request,
        "ventes/vente_detail.html",
        {
            "vente": vente,
        },
    )


def vente_create(request):
    if request.method == "POST":
        vente = Vente.objects.create()
        return redirect("vente_detail", vente.id)

    medicaments = Medicament.objects.select_related('catalogue').filter(stock__gt=0)
    clients = Client.objects.all()

    return render(request, "ventes/vente_form.html", {
        "medicaments": medicaments,
        "clients": clients,
    })


@csrf_exempt
@require_POST
def vente_directe_save(request):
    try:
        data = json.loads(request.body)

        if "lignes" not in data:
            return JsonResponse({"error": "Aucune ligne envoyée"}, status=400)

        with transaction.atomic():
            vente = Vente.objects.create(total=0)

            if data.get('client_id'):
                try:
                    client = Client.objects.get(id=data['client_id'])
                    vente.client = client
                    vente.save()
                except Client.DoesNotExist:
                    return JsonResponse({"error": "Client introuvable"}, status=400)

            total = 0

            for ligne in data["lignes"]:
                try:
                    med = Medicament.objects.get(id=ligne["medicament_id"])
                except Medicament.DoesNotExist:
                    return JsonResponse({"error": "Médicament introuvable"}, status=400)

                qte = int(ligne.get("quantite", 0))

                if qte <= 0:
                    return JsonResponse({"error": "Quantité invalide"}, status=400)

                if med.stock < qte:
                    return JsonResponse(
                        {"error": f"Stock insuffisant pour {med.catalogue.nom}"},
                        status=400
                    )

                montant = qte * med.prix

                LigneVente.objects.create(
                    vente=vente,
                    medicament=med,
                    quantite=qte,
                    prix_unitaire=med.prix,
                    montant_total=montant
                )

                med.stock -= qte
                med.save()

                total += montant

            vente.total = total
            vente.save()

        return JsonResponse({
            "success": True,
            "id": vente.id,
            "total": total
        })

    except json.JSONDecodeError:
        return JsonResponse({"error": "JSON invalide"}, status=400)

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


# ===================== API FLUTTER =====================

def _serialize_ligne(l):
    return {
        "medicament": l.medicament.catalogue.nom if l.medicament and l.medicament.catalogue else str(l.medicament),
        "quantite": l.quantite,
        "prix_unitaire": float(l.prix_unitaire or 0),
        "montant_total": float(l.montant_total or 0),
    }


def _serialize_vente(v):
    try:
        if v.ordonnance and v.ordonnance.consultation:
            client = v.ordonnance.consultation.client.nom
        else:
            client = v.client.nom if v.client else "—"
    except Exception:
        client = "—"

    return {
        "id": v.id,
        "date": v.date.strftime("%d/%m/%Y %H:%M"),
        "client": client,
        "total": float(v.total or 0),
        "nb_lignes": v.lignes.count(),
        "lignes": [_serialize_ligne(l) for l in v.lignes.select_related("medicament__catalogue").all()],
    }


def api_ventes_liste(request):
    """GET /ventes/api/liste/ — 50 dernières ventes."""
    ventes = (
        Vente.objects
        .select_related("client", "ordonnance__consultation__client")
        .prefetch_related("lignes__medicament__catalogue")
        .order_by("-date")[:50]
    )
    data = [_serialize_vente(v) for v in ventes]
    return JsonResponse(data, safe=False)


def api_ventes_stats(request):
    """GET /ventes/api/stats/"""
    today = now()
    ventes = Vente.objects.all()

    total_general = ventes.aggregate(Sum("total"))["total__sum"] or 0
    total_jour = ventes.filter(
        date__year=today.year, date__month=today.month, date__day=today.day,
    ).aggregate(Sum("total"))["total__sum"] or 0
    nb_ventes = ventes.count()
    ticket_moyen = round(total_general / nb_ventes, 0) if nb_ventes > 0 else 0

    return JsonResponse({
        "total_general": float(total_general),
        "total_jour": float(total_jour),
        "nb_ventes": nb_ventes,
        "ticket_moyen": float(ticket_moyen),
    })


def api_vente_jour(request):
    """
    GET /ventes/api/vente-jour/?date=YYYY-MM-DD
    Détail complet des ventes d'une journée (par défaut aujourd'hui).
    """
    date_str = request.GET.get("date")
    if date_str:
        jour = date_str
        ventes = Vente.objects.filter(date__date=jour)
    else:
        today = now()
        jour = today.strftime("%Y-%m-%d")
        ventes = Vente.objects.filter(
            date__year=today.year, date__month=today.month, date__day=today.day,
        )

    ventes = (
        ventes
        .select_related("client", "ordonnance__consultation__client")
        .prefetch_related("lignes__medicament__catalogue")
        .order_by("-date")
    )

    total = ventes.aggregate(Sum("total"))["total__sum"] or 0
    nb_ventes = ventes.count()

    produits = {}
    for v in ventes:
        for l in v.lignes.all():
            nom = l.medicament.catalogue.nom if l.medicament and l.medicament.catalogue else "Inconnu"
            if nom not in produits:
                produits[nom] = {"nom": nom, "quantite": 0, "montant": 0.0}
            produits[nom]["quantite"] += l.quantite
            produits[nom]["montant"] += float(l.montant_total or 0)

    top_produits = sorted(produits.values(), key=lambda p: p["montant"], reverse=True)

    return JsonResponse({
        "date": jour,
        "total": float(total),
        "nb_ventes": nb_ventes,
        "ticket_moyen": float(round(total / nb_ventes, 0)) if nb_ventes > 0 else 0,
        "ventes": [_serialize_vente(v) for v in ventes],
        "top_produits": top_produits,
    })


# ===================== EXPORTS VENTE DU JOUR =====================

def export_vente_jour_excel(request):
    """GET /ventes/export/jour/excel/?date=YYYY-MM-DD"""
    date_str = request.GET.get("date")
    if date_str:
        jour_label = date_str
        ventes = Vente.objects.filter(date__date=date_str)
    else:
        today = now()
        jour_label = today.strftime("%Y-%m-%d")
        ventes = Vente.objects.filter(
            date__year=today.year, date__month=today.month, date__day=today.day,
        )

    ventes = (
        ventes
        .select_related("client", "ordonnance__consultation__client")
        .order_by("-date")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Vente du jour"

    headers = ["N°", "Heure", "Client", "Montant (FCFA)"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D4F", end_color="2E7D4F", fill_type="solid")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total = 0
    for v in ventes:
        try:
            if v.ordonnance and v.ordonnance.consultation:
                client = v.ordonnance.consultation.client.nom
            else:
                client = v.client.nom if v.client else "Anonyme"
        except Exception:
            client = "Anonyme"

        total += v.total or 0
        ws.append([v.id, v.date.strftime("%H:%M"), client, float(v.total or 0)])

    ws.append(["", "", "TOTAL", float(total)])
    last_row = ws.max_row
    for col in range(1, 5):
        ws.cell(row=last_row, column=col).font = Font(bold=True)

    for col_letter in ["A", "B", "C", "D"]:
        ws.column_dimensions[col_letter].width = 20

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="vente_du_jour_{jour_label}.xlsx"'
    wb.save(response)
    return response


def export_vente_jour_pdf(request):
    """GET /ventes/export/jour/pdf/?date=YYYY-MM-DD"""
    date_str = request.GET.get("date")
    if date_str:
        jour_label = date_str
        ventes = Vente.objects.filter(date__date=date_str)
    else:
        today = now()
        jour_label = today.strftime("%d/%m/%Y")
        ventes = Vente.objects.filter(
            date__year=today.year, date__month=today.month, date__day=today.day,
        )

    ventes = (
        ventes
        .select_related("client", "ordonnance__consultation__client")
        .order_by("-date")
    )

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="vente_du_jour.pdf"'

    doc = SimpleDocTemplate(
        response, pagesize=A4,
        rightMargin=1.2 * cm, leftMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )

    styles = getSampleStyleSheet()
    elements = []

    titre_style = styles["Heading2"]
    titre_style.alignment = TA_CENTER
    elements.append(Paragraph("<b>VENTE DU JOUR — PARCELLES VÉTO</b>", titre_style))
    elements.append(Paragraph(
        f"<para alignment='center'><font size='10'>{jour_label}</font></para>",
        styles["Normal"],
    ))
    elements.append(Spacer(1, 0.5 * cm))

    data = [["N°", "Heure", "Client", "Montant (FCFA)"]]
    total = 0
    for v in ventes:
        try:
            if v.ordonnance and v.ordonnance.consultation:
                client = v.ordonnance.consultation.client.nom
            else:
                client = v.client.nom if v.client else "Anonyme"
        except Exception:
            client = "Anonyme"

        total += v.total or 0
        data.append([str(v.id), v.date.strftime("%H:%M"), client, f"{v.total:,.0f}"])

    data.append(["", "", "TOTAL", f"{total:,.0f} FCFA"])

    table = Table(data, colWidths=[2*cm, 3*cm, 7*cm, 4*cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D4F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -2), colors.whitesmoke),
        ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -2), 9),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F5E9")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.darkgreen),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.6 * cm))

    resume_style = styles["Normal"]
    resume_style.alignment = TA_RIGHT
    elements.append(Paragraph(f"<b>Nombre de ventes : {ventes.count()}</b>", resume_style))

    doc.build(elements)
    return response