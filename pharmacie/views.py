from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.http import require_GET
from django.utils import timezone
from django.http import JsonResponse, HttpResponse
from django.db.models import Q, F, Sum
from datetime import timedelta
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
import json
from django.views.decorators.http import require_POST
from django.contrib.auth.decorators import login_required
import logging

logger = logging.getLogger(__name__)

# ===================== EXPORTS (Excel / PDF) =====================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from django.urls import reverse
from django.db import transaction

# ===================== MODELS IMPORTS =====================
from .models import Medicament, CatalogueMedicament, FamilleMedicament
from fournisseurs.models import Fournisseur
from ventes.models import Vente, LigneVente
from consultations.models import Consultation, RendezVous 
from clients.models import Client
from animaux.models import Animal

# ===================== NOTIFICATIONS FCM =====================
from notifications.firebase_utils import notify_all_docteurs

# ===================== DRF =====================
from rest_framework import viewsets
from .serializers import MedicamentSerializer, FamilleMedicamentSerializer


# ===================== DRF VIEWSETS =====================
class FamilleMedicamentViewSet(viewsets.ModelViewSet):
    queryset = FamilleMedicament.objects.all()
    serializer_class = FamilleMedicamentSerializer


class MedicamentViewSet(viewsets.ModelViewSet):
    queryset = Medicament.objects.select_related('catalogue__famille').all()
    serializer_class = MedicamentSerializer


# ===================== API VIEWS =====================
def api_medicaments(request):
    search = request.GET.get("search", "")
    famille = request.GET.get("famille", "")

    qs = Medicament.objects.select_related(
        "catalogue__famille",
        "fournisseur"
    ).all()

    if search:
        qs = qs.filter(
            Q(catalogue__nom__icontains=search) |
            Q(catalogue__famille__nom__icontains=search) |
            Q(fournisseur__nom__icontains=search)
        )

    if famille:
        qs = qs.filter(catalogue__famille__nom=famille)

    data = []
    for m in qs:
        nom = m.catalogue.nom if m.catalogue else "Médicament sans nom"
        famille_nom = m.catalogue.famille.nom if (m.catalogue and m.catalogue.famille) else ""

        data.append({
            "id": m.id,
            "nom": nom,
            "famille": famille_nom,
            "fournisseur": m.fournisseur.nom if m.fournisseur else "",
            "stock": m.stock,
            "prix": float(m.prix or 0),
            "seuil_alerte": m.seuil_alerte,
            "statut": (
                "Rupture" if m.stock == 0
                else "Alerte" if m.stock <= m.seuil_alerte
                else "OK"
            )
        })

    return JsonResponse({
        "count": len(data),
        "results": data
    })


@require_http_methods(["GET"])
def api_medicaments_liste(request):
    qs = Medicament.objects.select_related(
        "catalogue__famille", "fournisseur"
    ).all().order_by("catalogue__nom")

    search = request.GET.get("search", "")
    if search:
        qs = qs.filter(
            Q(catalogue__nom__icontains=search) |
            Q(catalogue__famille__nom__icontains=search)
        )

    data = []
    for m in qs:
        data.append({
            "id": m.id,
            "nom": m.catalogue.nom if m.catalogue else "",
            "famille": m.catalogue.famille.nom if m.catalogue and m.catalogue.famille else "",
            "fournisseur": m.fournisseur.nom if m.fournisseur else "",
            "fournisseur_id": m.fournisseur.id if m.fournisseur else None,
            "stock": m.stock,
            "prix": float(m.prix or 0),
            "seuil_alerte": m.seuil_alerte,
            "statut": (
                "Rupture" if m.stock == 0
                else "Alerte" if m.stock <= m.seuil_alerte
                else "OK"
            ),
        })

    return JsonResponse(data, safe=False)


# ===================== DJANGO VIEWS =====================

def pharmacie_dashboard(request):
    today = timezone.localdate()

    prochains_rdv = (
        RendezVous.objects
        .select_related("animal", "animal__client")
        .filter(date_rdv__date__gte=today)
        .order_by("date_rdv")
    )

    prochains_rdv = sorted(
        prochains_rdv,
        key=lambda rdv: (
            rdv.date_rdv.date() != today,
            rdv.date_rdv.date(),
            rdv.date_rdv.time()
        )
    )

    total_consultations = Consultation.objects.count()
    total_rdv = RendezVous.objects.count()

    recettes_aujourdhui = Vente.objects.filter(
        date__date=today
    ).aggregate(total=Sum("total"))["total"] or 0

    total_medicaments = Medicament.objects.count()
    total_fournisseurs = Fournisseur.objects.count() if 'Fournisseur' in globals() else 0
    
    stock_critique = Medicament.objects.filter(
        stock__lte=F('seuil_alerte')
    ).select_related('catalogue', 'catalogue__famille')

    medicaments_recent = Medicament.objects.select_related('catalogue', 'catalogue__famille').order_by('-id')[:5]
    dernieres_ventes = Vente.objects.select_related(
        'ordonnance__consultation__client', 
        'client'
    ).order_by('-date')[:5]

    jours_mois = []
    ventes_mois = []
    consultations_mois = []

    for i in range(6, -1, -1):
        day = today - timedelta(days=i)
        jours_mois.append(day.strftime("%d/%m"))

        ventes_jour = Vente.objects.filter(date__date=day).aggregate(
            total=Sum("total")
        )["total"] or 0

        consult_jour = Consultation.objects.filter(date__date=day).count()

        ventes_mois.append(ventes_jour)
        consultations_mois.append(consult_jour)

    context = {
        "today": today,
        "prochains_rdv": prochains_rdv,
        "total_consultations": total_consultations,
        "total_rdv": total_rdv,
        "recettes_aujourdhui": recettes_aujourdhui,
        "total_medicaments": total_medicaments,
        "total_fournisseurs": total_fournisseurs,
        "stock_critique": stock_critique,
        "medicaments_recent": medicaments_recent,
        "dernieres_ventes": dernieres_ventes,
        "jours_mois": jours_mois,
        "ventes_mois": ventes_mois,
        "consultations_mois": consultations_mois,
    }

    return render(request, "pharmacie/dashboard.html", context)


def medicaments_list(request):
    search = request.GET.get("search", "")
    famille_id = request.GET.get("famille", "")

    medicaments = Medicament.objects.select_related(
        'catalogue__famille', 'fournisseur'
    ).all().order_by("catalogue__nom")

    if search:
        medicaments = medicaments.filter(
            Q(catalogue__nom__icontains=search) |
            Q(catalogue__famille__nom__icontains=search) |
            Q(fournisseur__nom__icontains=search)
        )

    if famille_id:
        medicaments = medicaments.filter(catalogue__famille_id=famille_id)

    context = {
        "medicaments": medicaments,
        "familles": FamilleMedicament.objects.all(),
        "search": search,
    }

    return render(request, "pharmacie/medicaments_list.html", context)


def medicament_create(request):
    if request.method == "POST":
        catalogue_id = request.POST.get("catalogue")
        stock = int(request.POST.get("stock", 0))
        prix = request.POST.get("prix")
        seuil = request.POST.get("seuil_alerte", 5)

        fournisseur_nom = request.POST.get("fournisseur_nom")
        fournisseur_contact = request.POST.get("fournisseur_contact")
        fournisseur_email = request.POST.get("fournisseur_email")
        fournisseur_adresse = request.POST.get("fournisseur_adresse")

        catalogue = get_object_or_404(CatalogueMedicament, id=catalogue_id)

        fournisseur = None
        if fournisseur_nom:
            fournisseur, created = Fournisseur.objects.get_or_create(
                nom=fournisseur_nom,
                defaults={
                    "telephone": fournisseur_contact,
                    "email": fournisseur_email,
                    "adresse": fournisseur_adresse,
                }
            )
            if not created:
                fournisseur.telephone = fournisseur_contact
                fournisseur.email = fournisseur_email
                fournisseur.adresse = fournisseur_adresse
                fournisseur.save()

        medicament, created = Medicament.objects.get_or_create(
            catalogue=catalogue,
            defaults={
                "stock": stock,
                "prix": prix,
                "seuil_alerte": seuil,
                "fournisseur": fournisseur
            }
        )
        if not created:
            medicament.stock += stock
            medicament.prix = prix
            medicament.seuil_alerte = seuil
            medicament.fournisseur = fournisseur
            medicament.save()

        return redirect("medicaments_list")

    catalogues = CatalogueMedicament.objects.all()
    familles = FamilleMedicament.objects.all()

    return render(request, "pharmacie/medicament_form.html", {
        "catalogues": catalogues,
        "familles": familles,
    })


@csrf_exempt
def dashboard_stats(request):
    today = timezone.localdate()
    data = {
        "clients": Client.objects.count(),
        "animaux": Animal.objects.count(),
        "consultations": Consultation.objects.count(),
        "rendezvous": RendezVous.objects.count(),
        "medicaments": Medicament.objects.count(),
        "stock_alertes": Medicament.objects.filter(
            stock__lte=F("seuil_alerte")
        ).count(),
        "ventes_jour": float(
            Vente.objects.filter(date__date=today)
            .aggregate(total=Sum("total"))["total"] or 0
        ),
    }
    return JsonResponse(data)


@require_GET
def api_stock_alertes(request):
    alertes = Medicament.objects.filter(
        stock__lte=F("seuil_alerte")
    ).select_related("catalogue", "fournisseur")

    data = [{
        "medicament": m.catalogue.nom,
        "stock": m.stock,
        "seuil": m.seuil_alerte,
        "fournisseur": m.fournisseur.nom if m.fournisseur else "",
    } for m in alertes]

    return JsonResponse(data, safe=False)


@require_GET
def api_fournisseurs(request):
    fournisseurs = Fournisseur.objects.all()

    data = [{
        "id": f.id,
        "nom": f.nom,
        "contact": f.telephone or "",
        "email": f.email or "",
        "adresse": f.adresse or "",
    } for f in fournisseurs]

    return JsonResponse(data, safe=False)


def api_alertes(request):
    alertes = (
        Medicament.objects
        .select_related("catalogue", "fournisseur")
        .filter(stock__lte=F("seuil_alerte"))
        .order_by("stock")
    )
    data = []
    for m in alertes:
        data.append({
            "medicament": m.catalogue.nom if m.catalogue else str(m),
            "stock": m.stock,
            "seuil": m.seuil_alerte,
            "fournisseur": m.fournisseur.nom if m.fournisseur else "",
        })
    return JsonResponse(data, safe=False)


@csrf_exempt
@require_http_methods(["POST"])
def api_ajouter_medicament(request):
    try:
        data = json.loads(request.body)

        famille_nom = data.get("famille", "Général")
        famille, _ = FamilleMedicament.objects.get_or_create(nom=famille_nom)

        catalogue_nom = data.get("nom", "")
        if not catalogue_nom:
            return JsonResponse({"error": "Nom obligatoire"}, status=400)

        catalogue, _ = CatalogueMedicament.objects.get_or_create(
            nom=catalogue_nom,
            defaults={"famille": famille}
        )

        fournisseur = None
        fournisseur_id = data.get("fournisseur_id")
        if fournisseur_id:
            try:
                fournisseur = Fournisseur.objects.get(pk=fournisseur_id)
            except Exception:
                pass

        medicament = Medicament.objects.create(
            catalogue=catalogue,
            stock=int(data.get("stock", 0)),
            prix=float(data.get("prix", 0)),
            seuil_alerte=int(data.get("seuil_alerte", 5)),
            fournisseur=fournisseur,
        )

        return JsonResponse({
            "id": medicament.id,
            "nom": medicament.catalogue.nom,
            "stock": medicament.stock,
        }, status=201)

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=400)


@csrf_exempt
@require_http_methods(["PUT"])
def api_modifier_medicament(request, medicament_id):
    try:
        medicament = Medicament.objects.select_related(
            "catalogue__famille", "fournisseur"
        ).get(pk=medicament_id)

        data = json.loads(request.body)

        if "stock" in data:
            medicament.stock = int(data["stock"])
        if "prix" in data:
            medicament.prix = float(data["prix"])
        if "seuil_alerte" in data:
            medicament.seuil_alerte = int(data["seuil_alerte"])

        fournisseur_id = data.get("fournisseur_id")
        if fournisseur_id:
            try:
                medicament.fournisseur = Fournisseur.objects.get(pk=fournisseur_id)
            except Exception:
                pass

        if "nom" in data and data["nom"]:
            medicament.catalogue.nom = data["nom"]
            medicament.catalogue.save()

        if "famille" in data and data["famille"]:
            famille, _ = FamilleMedicament.objects.get_or_create(nom=data["famille"])
            medicament.catalogue.famille = famille
            medicament.catalogue.save()

        medicament.save()

        return JsonResponse({
            "id": medicament.id,
            "nom": medicament.catalogue.nom,
            "stock": medicament.stock,
        })

    except Medicament.DoesNotExist:
        return JsonResponse({"error": "Médicament introuvable"}, status=404)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=400)


@csrf_exempt
@require_http_methods(["DELETE"])
def api_supprimer_medicament(request, medicament_id):
    try:
        medicament = Medicament.objects.get(pk=medicament_id)
        medicament.delete()
        return JsonResponse({"success": True})
    except Medicament.DoesNotExist:
        return JsonResponse({"error": "Médicament introuvable"}, status=404)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=400)


@csrf_exempt
@require_POST
def vente_directe_save(request):
    try:
        data = json.loads(request.body)

        if "lignes" not in data:
            return JsonResponse({"error": "Aucune ligne envoyée"}, status=400)

        with transaction.atomic():
            vente = Vente.objects.create(total=0)

            if data.get('client_id'):
                try:
                    client = Client.objects.get(id=data['client_id'])
                    vente.client = client
                    vente.save()
                except Client.DoesNotExist:
                    return JsonResponse({"error": "Client introuvable"}, status=400)

            total = 0

            for ligne in data["lignes"]:
                try:
                    med = Medicament.objects.select_for_update().get(id=ligne["medicament_id"])
                except Medicament.DoesNotExist:
                    return JsonResponse({"error": "Médicament introuvable"}, status=400)

                qte = int(ligne.get("quantite", 0))

                if qte <= 0:
                    return JsonResponse({"error": "Quantité invalide"}, status=400)

                if med.stock < qte:
                    return JsonResponse(
                        {"error": f"Stock insuffisant pour {med.catalogue.nom}"},
                        status=400
                    )

                montant = qte * med.prix

                LigneVente.objects.create(
                    vente=vente,
                    medicament=med,
                    quantite=qte,
                    prix_unitaire=med.prix,
                    montant_total=montant
                )

                med.stock -= qte
                med.save()

                # 💡 Notification automatique si le stock devient critique ou bas
                if med.stock <= med.seuil_alerte:
                    nom_med = med.catalogue.nom if med.catalogue else "Médicament"
                    notify_all_docteurs(
                        title="⚠️ Alerte Stock Critique",
                        body=f"Stock bas pour '{nom_med}' (Restant: {med.stock}).",
                        data={"type": "stock", "medicament_id": str(med.id)}
                    )

                total += montant

            vente.total = total
            vente.save()

        # 💡 Notification globale pour la validation de la vente
        notify_all_docteurs(
            title="💰 Nouvelle Vente Validée",
            body=f"Une vente d'un montant de {total} FCFA a été enregistrée.",
            data={"type": "vente", "id": str(vente.id)}
        )

        return JsonResponse({
            "success": True,
            "id": vente.id,
            "total": total
        })

    except json.JSONDecodeError:
        return JsonResponse({"error": "JSON invalide"}, status=400)

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


@login_required
@require_POST
def terminer_consultation(request, consultation_id):
    consultation = get_object_or_404(Consultation, id=consultation_id)

    if consultation.statut == "terminee":
        return JsonResponse({
            'status': 'OK',
            'message': 'Consultation déjà terminée.',
            'redirect_url': reverse("consultation_detail", kwargs={'consultation_id': consultation.id})
        })

    ordonnance = Ordonnance.objects.filter(consultation=consultation).first()

    if not ordonnance:
        return JsonResponse({
            'error': "Aucune ordonnance n'a été trouvée pour cette consultation.",
            'redirect_url': reverse("ordonnance_create", kwargs={'consultation_id': consultation.id})
        }, status=400)

    if not ordonnance.lignes.exists():
        return JsonResponse({
            'error': "Impossible de terminer la consultation : l'ordonnance doit contenir au moins un médicament."
        }, status=400)

    try:
        with transaction.atomic():
            for ligne in ordonnance.lignes.select_related("medicament").select_for_update():
                if ligne.medicament.stock < ligne.quantite:
                    return JsonResponse({
                        'error': f"Stock insuffisant pour {ligne.medicament.catalogue.nom}. Stock actuel : {ligne.medicament.stock}"
                    }, status=400)

            vente = Vente.objects.create(
                ordonnance=ordonnance,
                total=0
            )

            total = 0
            for ligne in ordonnance.lignes.all():
                montant = ligne.quantite * ligne.medicament.prix
                LigneVente.objects.create(
                    vente=vente,
                    medicament=ligne.medicament,
                    quantite=ligne.quantite,
                    prix_unitaire=ligne.medicament.prix,
                    montant_total=montant
                )
                total += montant

                med = ligne.medicament
                med.stock -= ligne.quantite
                med.save()

                # 💡 Alerte stock critique si besoin après clôture de consultation
                if med.stock <= med.seuil_alerte:
                    nom_med = med.catalogue.nom if med.catalogue else "Médicament"
                    notify_all_docteurs(
                        title="⚠️ Alerte Stock Critique",
                        body=f"Stock bas pour '{nom_med}' (Restant: {med.stock}).",
                        data={"type": "stock", "medicament_id": str(med.id)}
                    )

            vente.total = total
            vente.save()

            consultation.statut = "terminee"
            consultation.save()

        # 💡 Notification de fin de consultation
        nom_animal = consultation.animal.nom if consultation.animal else "Patient"
        notify_all_docteurs(
            title="🩺 Consultation Clôturée",
            body=f"La consultation pour {nom_animal} a été finalisée.",
            data={"type": "consultation", "id": str(consultation.id)}
        )

        return JsonResponse({
            'status': 'OK',
            'message': 'Consultation terminée avec succès.',
            'redirect_url': reverse("vente_detail", kwargs={'vente_id': vente.id})
        })

    except Exception as exc:
        logger.error(f"Erreur terminer_consultation: {exc}")
        return JsonResponse({
            'error': "Une erreur est survenue lors de la validation de la consultation."
        }, status=500)


def export_pharmacie_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pharmacie"

    headers = ["Médicament", "Famille", "Stock", "Seuil alerte", "Prix (FCFA)", "Statut", "Fournisseur"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D4F", end_color="2E7D4F", fill_type="solid")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    medicaments = Medicament.objects.select_related("catalogue__famille", "fournisseur").order_by("catalogue__nom")

    for m in medicaments:
        statut = "Rupture" if m.stock == 0 else "Alerte" if m.stock <= m.seuil_alerte else "OK"
        ws.append([
            m.catalogue.nom if m.catalogue else "",
            m.catalogue.famille.nom if m.catalogue and m.catalogue.famille else "",
            m.stock,
            m.seuil_alerte,
            float(m.prix or 0),
            statut,
            m.fournisseur.nom if m.fournisseur else "",
        ])

    for col_letter in ["A", "B", "C", "D", "E", "F", "G"]:
        ws.column_dimensions[col_letter].width = 20

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="pharmacie.xlsx"'
    wb.save(response)
    return response


def export_pharmacie_pdf(request):
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="pharmacie.pdf"'

    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        rightMargin=1.2 * cm, leftMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )

    styles = getSampleStyleSheet()
    elements = []

    titre_style = styles["Heading2"]
    titre_style.alignment = TA_CENTER
    elements.append(Paragraph("**ÉTAT DU STOCK - PHARMACIE PARCELLES VÉTO**", titre_style))
    elements.append(Paragraph(
        f"Édité le {timezone.now().strftime('%d/%m/%Y %H:%M')}",
        styles["Normal"],
    ))
    elements.append(Spacer(1, 0.5 * cm))

    data = [["Médicament", "Famille", "Stock", "Seuil", "Prix (FCFA)", "Statut", "Fournisseur"]]
    medicaments = Medicament.objects.select_related("catalogue__famille", "fournisseur").order_by("catalogue__nom")

    for m in medicaments:
        statut = "Rupture" if m.stock == 0 else "Alerte" if m.stock <= m.seuil_alerte else "OK"
        data.append([
            m.catalogue.nom if m.catalogue else "",
            m.catalogue.famille.nom if m.catalogue and m.catalogue.famille else "",
            str(m.stock),
            str(m.seuil_alerte),
            f"{m.prix:,.0f}",
            statut,
            m.fournisseur.nom if m.fournisseur else "-",
        ])

    table = Table(data, colWidths=[4*cm, 3*cm, 2*cm, 2*cm, 3*cm, 2.5*cm, 4*cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D4F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
    ]))
    elements.append(table)

    doc.build(elements)
    return response


@login_required
@require_POST
def ajouter_famille_ajax(request):
    nom = request.POST.get("nom", "").strip()
    if not nom:
        return JsonResponse({"success": False, "error": "Le nom de la famille est obligatoire."}, status=400)
    
    famille, created = FamilleMedicament.objects.get_or_create(nom=nom)
    return JsonResponse({
        "success": True,
        "id": famille.id,
        "nom": famille.nom,
        "created": created
    })


@login_required
@require_POST
def ajouter_catalogue_ajax(request):
    nom = request.POST.get("nom", "").strip()
    famille_id = request.POST.get("famille_id")

    if not nom or not famille_id:
        return JsonResponse({"success": False, "error": "Tous les champs sont requis."}, status=400)

    try:
        famille = FamilleMedicament.objects.get(id=famille_id)
        med, created = CatalogueMedicament.objects.get_or_create(
            nom=nom,
            famille=famille
        )
        return JsonResponse({
            "success": True,
            "id": med.id,
            "nom": med.nom
        })
    except FamilleMedicament.DoesNotExist:
        return JsonResponse({"success": False, "error": "Famille introuvable."}, status=404)
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@csrf_exempt
@login_required
@require_POST
def api_creer_medicament_express(request):
    try:
        data = json.loads(request.body)
        nom = data.get("nom", "").strip()
        famille_nom = data.get("famille", "Général").strip()
        stock = int(data.get("stock", 0))
        prix = float(data.get("prix", 0))
        seuil = int(data.get("seuil_alerte", 5))

        if not nom:
            return JsonResponse({"success": False, "error": "Le nom du médicament est requis."}, status=400)

        famille, _ = FamilleMedicament.objects.get_or_create(nom=famille_nom)

        catalogue, _ = CatalogueMedicament.objects.get_or_create(
            nom=nom,
            defaults={"famille": famille}
        )

        medicament, created = Medicament.objects.get_or_create(
            catalogue=catalogue,
            defaults={
                "stock": stock,
                "prix": prix,
                "seuil_alerte": seuil,
            }
        )

        if not created:
            medicament.stock += stock
            medicament.prix = prix
            medicament.save()

        return JsonResponse({
            "success": True,
            "medicament": {
                "id": medicament.id,
                "nom": medicament.catalogue.nom,
                "prix": float(medicament.prix),
                "stock": medicament.stock,
            }
        }, status=201)

    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)


@login_required
def recherche_rapide_medicament(request):
    query = request.GET.get('q', '').strip()
    resultats = []
    
    if len(query) >= 2: 
        medicaments = Medicament.objects.filter(
            Q(catalogue__nom__icontains=query) |
            Q(catalogue__famille__nom__icontains=query)
        ).select_related('catalogue', 'catalogue__famille')[:8]
        
        for med in medicaments:
            resultats.append({
                'id': med.id,
                'nom': med.catalogue.nom,
                'famille': med.catalogue.famille.nom if med.catalogue.famille else 'Générique',
                'prix': f"{med.prix:.0f} FCFA",
                'stock': med.stock,
                'en_alerte': med.stock <= med.seuil_alerte,
                'en_rupture': med.stock == 0,
            })
            
    return JsonResponse({'medicaments': resultats}) 


def recherche_medicament_page(request):
    query = request.GET.get('q', '')
    familles = FamilleMedicament.objects.all()
    
    return render(request, 'pharmacie/medicaments_list.html', {
        'familles': familles,
        'query': query,
    })


def medicament_detail(request, pk):
    medicament = get_object_or_404(Medicament, pk=pk)
    return render(request, 'pharmacie/medicament_detail.html', {'medicament': medicament})