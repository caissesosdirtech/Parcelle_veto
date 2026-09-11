import json
from django.shortcuts import render
from django.db.models import Sum
from django.db.models.functions import TruncDate
from django.utils import timezone
from django.http import HttpResponse, JsonResponse
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied

# Dépendances ReportLab (PDF)
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

# Dépendances OpenPyXL (Excel)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Modèles
from ventes.models import Vente


# ── VUE WEB ───────────────────────────────────────────────────────────────────

@login_required
def rapport_caisse(request):
    """Affiche le rapport de caisse avec statistiques et historique filtrable."""
    if hasattr(request.user, 'role') and request.user.role == 'EMPLOYE' and not request.user.is_superuser:
        raise PermissionDenied("Accès refusé : Le rapport de caisse est strictly réservé aux docteurs.")

    date_debut = request.GET.get("date_debut")
    date_fin = request.GET.get("date_fin")

    ventes = (
        Vente.objects.select_related(
            "client",
            "ordonnance__consultation__client",
            "ordonnance__consultation__animal",
        )
        .prefetch_related("lignes__medicament")
        .order_by("-date", "-id")
    )

    if date_debut:
        ventes = ventes.filter(date__date__gte=date_debut)
    if date_fin:
        ventes = ventes.filter(date__date__lte=date_fin)

    total_recettes = ventes.aggregate(total=Sum("total"))["total"] or 0
    recettes_jour = Vente.objects.filter(
        date__date=timezone.now().date()
    ).aggregate(total=Sum("total"))["total"] or 0
    nombre_ventes = ventes.count()
    ticket_moyen = (total_recettes / nombre_ventes) if nombre_ventes > 0 else 0

    recettes = (
        ventes
        .annotate(jour=TruncDate("date"))
        .values("jour")
        .annotate(total=Sum("total"))
        .order_by("jour")
    )

    return render(request, "caisse/rapport_caisse.html", {
        "ventes": ventes,
        "total_recettes": total_recettes,
        "recettes_jour": recettes_jour,
        "nombre_ventes": nombre_ventes,
        "ticket_moyen": ticket_moyen,
        "recettes": recettes,
        "date_debut": date_debut,
        "date_fin": date_fin,
        "active_page": "caisse",
    })


# ── API FLUTTER ───────────────────────────────────────────────────────────────

def api_rapport_caisse(request):
    """
    GET /caisse/api/rapport/?date_debut=YYYY-MM-DD&date_fin=YYYY-MM-DD
    Retourne les statistiques et la liste des ventes pour CaisseScreen Flutter.
    """
    date_debut = request.GET.get("date_debut")
    date_fin = request.GET.get("date_fin")

    ventes_qs = (
        Vente.objects.select_related(
            "client",
            "ordonnance__consultation__client",
            "ordonnance__consultation__animal",
        )
        .order_by("-date", "-id")
    )

    if date_debut:
        ventes_qs = ventes_qs.filter(date__date__gte=date_debut)
    if date_fin:
        ventes_qs = ventes_qs.filter(date__date__lte=date_fin)

    total_recettes = ventes_qs.aggregate(total=Sum("total"))["total"] or 0
    recettes_jour = Vente.objects.filter(
        date__date=timezone.now().date()
    ).aggregate(total=Sum("total"))["total"] or 0
    nombre_ventes = ventes_qs.count()
    ticket_moyen = (float(total_recettes) / nombre_ventes) if nombre_ventes > 0 else 0

    ventes_data = []
    for v in ventes_qs:
        try:
            if v.ordonnance and v.ordonnance.consultation:
                client = v.ordonnance.consultation.client.nom
                animal = v.ordonnance.consultation.animal.nom
                type_vente = "Ordonnance"
            else:
                client = v.client.nom if v.client else "Anonyme"
                animal = ""
                type_vente = "Directe"
        except Exception:
            client = "Anonyme"
            animal = ""
            type_vente = "Directe"

        ventes_data.append({
            "id": v.id,
            "date": v.date.strftime("%d/%m/%Y"),
            "client": client,
            "animal": animal,
            "type": type_vente,
            "montant": float(v.total or 0),
        })

    return JsonResponse({
        "stats": {
            "total_recettes": float(total_recettes),
            "recettes_jour": float(recettes_jour),
            "nombre_ventes": nombre_ventes,
            "ticket_moyen": round(ticket_moyen, 0),
        },
        "ventes": ventes_data,
    }, status=200)


# ── PDF ───────────────────────────────────────────────────────────────────────

def rapport_caisse_pdf(request):
    """Génère un rapport de caisse imprimable au format PDF."""
    date_debut = request.GET.get("date_debut")
    date_fin = request.GET.get("date_fin")

    ventes = (
        Vente.objects.select_related(
            "client",
            "ordonnance__consultation__client",
            "ordonnance__consultation__animal",
        )
        .order_by("-date", "-id")
    )

    if date_debut:
        ventes = ventes.filter(date__date__gte=date_debut)
    if date_fin:
        ventes = ventes.filter(date__date__lte=date_fin)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="rapport_caisse_parcelles_veto.pdf"'

    doc = SimpleDocTemplate(
        response, pagesize=A4,
        rightMargin=1.2 * cm, leftMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # En-tête
    gauche = Paragraph(
        "<b>Dr. Ibrahima Pierre GUISSE</b><br/>"
        "Parcelles Assainies<br/>Thiès, Sénégal<br/>"
        "En face des cimetières Keur Dago<br/>"
        "Tél : 221 775385729 / 768331623<br/>"
        "E-mail : parcelles-veto@gmail.com",
        styles["Normal"],
    )
    droite = Paragraph(
        f'<para alignment="right"><b>Date :</b><br/>'
        f'{timezone.now().strftime("%d/%m/%Y %H:%M")}</para>',
        styles["Normal"],
    )
    header = Table([[gauche, droite]], colWidths=[12 * cm, 6 * cm])
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(header)
    elements.append(Spacer(1, 0.4 * cm))

    # Titre
    titre_style = styles["Heading2"]
    titre_style.alignment = TA_CENTER
    elements.append(Paragraph("<b>RAPPORT DE CAISSE - PARCELLES VETO</b>", titre_style))
    elements.append(Spacer(1, 0.5 * cm))

    # Tableau des ventes
    data = [["N°", "Date", "Client", "Animal", "Type", "Montant (FCFA)"]]
    total_general = 0

    for vente in ventes:
        try:
            if vente.ordonnance and vente.ordonnance.consultation:
                client = vente.ordonnance.consultation.client.nom
                animal = vente.ordonnance.consultation.animal.nom
                type_vente = "Ordonnance"
            else:
                client = vente.client.nom if vente.client else "Anonyme"
                animal = "-"
                type_vente = "Directe"
        except Exception:
            client = "Anonyme"
            animal = "-"
            type_vente = "Directe"

        montant = vente.total or 0
        total_general += montant
        data.append([
            str(vente.id),
            vente.date.strftime("%d/%m/%Y"),
            client, animal, type_vente,
            f"{montant:,.0f}",
        ])

    data.append(["", "", "", "", "TOTAL", f"{total_general:,.0f} FCFA"])

    table = Table(data, colWidths=[1.2 * cm, 3 * cm, 5 * cm, 4 * cm, 3 * cm, 3 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -2), colors.whitesmoke),
        ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -2), 9),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F5E9")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.darkgreen),
        ("FONTSIZE", (0, -1), (-1, -1), 10),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.6 * cm))

    resume_style = styles["Normal"]
    resume_style.alignment = TA_RIGHT
    elements.append(Paragraph(f"<b>Total général : {total_general:,.0f} FCFA</b>", resume_style))
    elements.append(Spacer(1, 1 * cm))
    elements.append(Paragraph(
        "<para alignment='center'><font size='9'>Parcelles Veto - La santé animale, notre priorité.</font></para>",
        styles["Normal"],
    ))

    doc.build(elements)
    return response


# ── EXPORT EXCEL ─────────────────────────────────────────────────────────────

def export_caisse_excel(request):
    """Exporte les ventes de la caisse sous format tableur Excel (.xlsx)."""
    date_debut = request.GET.get("date_debut")
    date_fin = request.GET.get("date_fin")

    ventes = (
        Vente.objects.select_related(
            "client",
            "ordonnance__consultation__client",
            "ordonnance__consultation__animal",
        )
        .order_by("-date", "-id")
    )

    if date_debut:
        ventes = ventes.filter(date__date__gte=date_debut)
    if date_fin:
        ventes = ventes.filter(date__date__lte=date_fin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport Caisse"

    headers = ["N°", "Date", "Client", "Animal", "Type", "Montant (FCFA)"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E7D4F", end_color="2E7D4F", fill_type="solid")
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total_general = 0
    for vente in ventes:
        try:
            if vente.ordonnance and vente.ordonnance.consultation:
                client = vente.ordonnance.consultation.client.nom
                animal = vente.ordonnance.consultation.animal.nom
                type_vente = "Ordonnance"
            else:
                client = vente.client.nom if vente.client else "Anonyme"
                animal = "-"
                type_vente = "Directe"
        except Exception:
            client = "Anonyme"
            animal = "-"
            type_vente = "Directe"

        montant = float(vente.total or 0)
        total_general += montant
        ws.append([
            vente.id,
            vente.date.strftime("%d/%m/%Y"),
            client, animal, type_vente,
            montant,
        ])

    # Ligne de total
    ws.append(["", "", "", "", "TOTAL", float(total_general)])
    last_row = ws.max_row
    for col in range(1, 7):
        ws.cell(row=last_row, column=col).font = Font(bold=True)

    for col_letter in ["A", "B", "C", "D", "E", "F"]:
        ws.column_dimensions[col_letter].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rapport_caisse.xlsx"'
    wb.save(response)
    return response